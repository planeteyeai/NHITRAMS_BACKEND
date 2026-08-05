"""Extract Inventory master data from RAMS Excel and write seed module."""
from __future__ import annotations

from collections import OrderedDict
from pathlib import Path

import openpyxl

EXCEL_PATH = Path(r"c:\Users\planeteye01\Downloads\782cf3c8bb674829a258b62ed4586925.xlsx")
OUT_PATH = Path(__file__).resolve().parent / "inventory" / "seed_inventory_masters.py"

DISPLAY_TO_REF = OrderedDict(
    [
        ("Road Furniture", "Road Furniture"),
        ("RCC Drain Line drain", "Drainage System"),
        ("Traffic Signs", "Traffic Signs"),
        ("Pavement Markings", "Pavement Markings"),
        ("Safety Assets", "Safety Assets"),
        ("Electrical Assets", "Electrical Assets"),
        ("Plantation", "Plantation"),
        ("Bus and Public Transport Facilities", "Bus and Public Transport Facilities"),
        ("Structures", "Structures"),
        ("Traffic Management Facilities", "Traffic Management Facilities"),
        ("Toll Plaza Facilities", "Toll Plaza Facilities"),
        ("Wayside Amenities", "Wayside Amenities"),
        ("Slope Protection", "Slope Protection"),
        ("Advanced Traffic Management System", "Advanced Traffic Management System (ATMS)"),
        ("Junctions", "Junctions"),
    ]
)

REF_HEADERS = set(DISPLAY_TO_REF.values())
NOISE = {
    "Project",
    "Distress Dimension",
    "Nos",
    "Width (w)",
    "Depth(d)",
    "Chainage (Auto Capture) KMZ file",
    "Type of Pavement",
}


def _normalize(name: str) -> str:
    return " ".join(str(name or "").split())


def extract_sides(ws) -> list[str]:
    sides = []
    for r in range(52, 78):
        v = ws.cell(r, 7).value
        if v is None:
            continue
        s = _normalize(v)
        if s and s not in sides:
            sides.append(s)
    return sides


def _parse_named_blocks(ws, col: int) -> dict[str, list[str]]:
    blocks: dict[str, list[str]] = {}
    current = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if v is None or not str(v).strip():
            continue
        name = _normalize(v)
        if name in NOISE:
            current = None
            continue
        if name in REF_HEADERS:
            current = name
            blocks.setdefault(current, [])
            continue
        if current is None:
            continue
        if name != current and name not in blocks[current]:
            blocks[current].append(name)
    return blocks


def extract_all(ws) -> OrderedDict[str, list[str]]:
    # Merge from multiple columns used in the Ref sheet.
    sources = [
        _parse_named_blocks(ws, 5),
        _parse_named_blocks(ws, 8),
        _parse_named_blocks(ws, 10),
        _parse_named_blocks(ws, 17),
    ]
    merged: OrderedDict[str, list[str]] = OrderedDict()
    for display, ref in DISPLAY_TO_REF.items():
        items: list[str] = []
        for src in sources:
            for it in src.get(ref, []):
                if it not in items and it != ref:
                    items.append(it)
        merged[display] = items
    return merged


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Ref_Inventory and Pavement "]
    sides = extract_sides(ws)
    cats = extract_all(ws)

    lines = [
        '"""Inventory category / side masters extracted from Ref_Inventory and Pavement Excel.',
        "",
        "Source of truth: RAMS APK workbook sheet `Ref_Inventory and Pavement`.",
        "Do not hand-edit dropdown values — re-run `_extract_inventory_masters.py` instead.",
        '"""',
        "",
        "INVENTORY_SIDE_SEED = [",
    ]
    for s in sides:
        lines.append(f"    {s!r},")
    lines.append("]")
    lines.append("")
    lines.append("INVENTORY_SEVERITY_SEED = ['Low', 'Medium', 'High']")
    lines.append("")
    lines.append("INVENTORY_CATEGORY_SEED = [")
    for i, cat in enumerate(cats, start=1):
        lines.append(f"    ({i}, {cat!r}),")
    lines.append("]")
    lines.append("")
    lines.append("INVENTORY_SUBCATEGORY_SEED = {")
    for cat, items in cats.items():
        lines.append(f"    {cat!r}: [")
        for it in items:
            lines.append(f"        {it!r},")
        lines.append("    ],")
    lines.append("}")
    lines.append("")

    OUT_PATH.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {OUT_PATH}")
    print("sides", len(sides), "categories", len(cats))
    for cat, items in cats.items():
        print(f"  {cat}: {len(items)}")


if __name__ == "__main__":
    main()

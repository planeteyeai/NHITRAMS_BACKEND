"""Generate seed_routine_maintenance_masters.py from Excel Ref_Routine Maintenance."""
from collections import OrderedDict
from pathlib import Path

import openpyxl

EXCEL = Path(r'c:\Users\planeteye01\Downloads\782cf3c8bb674829a258b62ed4586925.xlsx')
OUT = Path(r'c:\Users\planeteye01\Desktop\nhit\inventory\seed_routine_maintenance_masters.py')

# Category display order from Col C dropdown list (exact Excel text, stripped).
CATEGORY_ORDER = [
    'Median Maintenance',
    'Road Side Maintenance/ RoW (Avenue till boundary pillar)',
    'Maintenance and watering of Existing plants (Median + Avenue)',
    'Cleaning of Road furniture (Delineators, signboards, MBCB, PGR, Gantry Boards, KM, HM, 5th KM Stones, Guard post & Boundary pillars, Kerb etc.,)',
    'Cleaning of Road (MCW/SR)',
    'Cleaning of Drains',
    'Cleaning of Toll Plaza and Booths (from outside)',
    'Cleaning of waterway of structures/Structures',
    'Safety during service',
]

wb = openpyxl.load_workbook(EXCEL, data_only=True)
ws = wb['Ref_Routine Maintenance']

mapping = OrderedDict((name, []) for name in CATEGORY_ORDER)

for r in range(3, ws.max_row + 1):  # skip header row 2
    sno = ws.cell(r, 6).value
    head = ws.cell(r, 7).value
    particular = ws.cell(r, 8).value
    if head is None or particular is None:
        continue
    if isinstance(sno, str) and not str(sno).strip().isdigit():
        # section labels like Horticulture / Drainage
        continue
    cat = str(head).strip()
    sub = str(particular).strip()
    if not cat or not sub:
        continue
    if cat not in mapping:
        mapping[cat] = []
    if sub not in mapping[cat]:
        mapping[cat].append(sub)

# Verify all ordered categories present
missing = [c for c in CATEGORY_ORDER if c not in mapping or not mapping[c]]
if missing:
    raise SystemExit(f'Missing categories/subs: {missing}')

extra = [c for c in mapping if c not in CATEGORY_ORDER]
if extra:
    print('WARNING extra categories:', extra)

lines = [
    '"""Routine Maintenance masters from Ref_Routine Maintenance Excel (Annexure-I).',
    '',
    'Source of truth: RAMS APK workbook sheet `Ref_Routine Maintenance`.',
    'Keep long names exactly as in Excel — do not shorten.',
    'Re-run `_extract_rm_masters.py` / this generator to refresh.',
    '"""',
    '',
    'RM_CATEGORY_SEED = [',
]
for i, name in enumerate(CATEGORY_ORDER, start=1):
    lines.append(f'    ({i}, {name!r}),')
lines.append(']')
lines.append('')
lines.append('RM_SUBCATEGORY_SEED = {')
for name in CATEGORY_ORDER:
    lines.append(f'    {name!r}: [')
    for sub in mapping[name]:
        lines.append(f'        {sub!r},')
    lines.append('    ],')
lines.append('}')
lines.append('')
lines.append('# Side / severity are shared with inventory masters (see seed_inventory_masters).')
lines.append('')

OUT.write_text('\n'.join(lines), encoding='utf-8')
print(f'Wrote {OUT}')
total_subs = sum(len(v) for v in mapping.values())
print(f'Categories={len(CATEGORY_ORDER)} Subcategories={total_subs}')
for name in CATEGORY_ORDER:
    print(f'  - {name}: {len(mapping[name])} items')
    maxlen = max(len(s) for s in mapping[name])
    print(f'      longest sub={maxlen} chars; category len={len(name)}')
